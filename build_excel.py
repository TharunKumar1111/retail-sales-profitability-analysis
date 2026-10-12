import pandas as pd, numpy as np, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.properties import PageSetupProperties

F="Arial"; NAVY="1F3A5F"; TEAL="2A9D8F"; LIGHT="EAF2F8"; GREY="F4F6F7"
hfont=Font(name=F,bold=True,color="FFFFFF",size=10); hfill=PatternFill("solid",fgColor=NAVY)
bfont=Font(name=F,size=10); bold=Font(name=F,size=10,bold=True)
thin=Side(style="thin",color="D0D7DE"); box=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(ws,r,c,vals,fill=hfill):
    for i,v in enumerate(vals):
        x=ws.cell(r,c+i,v); x.font=hfont; x.fill=fill; x.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); x.border=box
def title(ws,r,text,sub=None):
    ws.cell(r,1,text).font=Font(name=F,size=12,bold=True,color=NAVY)
    if sub: ws.cell(r+1,1,sub).font=Font(name=F,size=9,italic=True,color="666666")

raw=pd.read_csv("data/retail_sales_raw.csv"); clean=pd.read_csv("data/retail_sales_clean.csv",parse_dates=["Order Date","Ship Date"])
log=pd.read_csv("data/cleaning_log.csv"); N=len(clean); last=N+1
wb=Workbook()

# ---------------- Read Me ----------------
ws=wb.active; ws.title="Read_Me"; ws.sheet_view.showGridLines=False
ws["A1"]="Retail Sales Performance & Profitability Analysis"; ws["A1"].font=Font(name=F,size=18,bold=True,color=NAVY)
ws["A2"]="Excel workbook  |  Project 1 of 2  |  Tools: Excel, SQL, Power BI"; ws["A2"].font=Font(name=F,size=10,italic=True,color="666666")
rows=[("Business context","An Indian multi-category retailer (Furniture, Office Supplies, Technology) sold across 5 regions from Jan-2023 to Dec-2025. Sales are growing, but profit margin is flat. Management wants to know where profit leaks and what to do about it."),
("Data","SYNTHETIC dataset (generated with a fixed random seed) modelled on retail order data. ~9,000 raw order lines with deliberate quality problems. Currency: INR."),
("Sheets","Raw_Data = untouched source (text dates, duplicates, messy text). Cleaning_Log = every fix and rows affected. Clean_Data = analysis-ready table + 7 formula helper columns. Summary = all analysis via SUMIFS/COUNTIFS/AVERAGEIFS. Customers = customer-level SUMIFS + rank. Dashboard = KPI cards and charts."),
("Excel skills shown","TRIM/PROPER/DATEVALUE cleaning, Remove Duplicates, named ranges, SUMIFS/COUNTIFS/AVERAGEIFS, INDEX-MATCH, RANK, nested IF, Excel Tables, conditional formatting, charts, Pivot Tables (build steps in docs)."),
("Colour legend","Blue font = hard-coded input/label that a person typed. Black font = formula. Red fill = loss-making. Everything on Summary/Customers/Dashboard recalculates if Clean_Data changes."),
("Assumptions","(1) Sales is net of discount (Sales = List Price x Qty x (1-Discount)). (2) Profit is after product cost. (3) Duplicates are defined as the same Order ID + Product ID. (4) Orders with Ship Date before Order Date or Quantity <= 0 were removed as invalid. (5) Missing City filled from the same customer's other orders. (6) Number of orders = count of first lines of each Order ID (Clean_Data sorted by Order ID)."),
("How to refresh","Paste new cleaned rows into Clean_Data (keep sorted by Order ID), extend the named ranges (Formulas > Name Manager), and everything else updates.")]
r=4
for k,v in rows:
    ws.cell(r,1,k).font=bold; ws.cell(r,1).fill=PatternFill("solid",fgColor=LIGHT); ws.cell(r,1).alignment=Alignment(vertical="top")
    c=ws.cell(r,2,v); c.font=bfont; c.alignment=Alignment(wrap_text=True,vertical="top"); ws.row_dimensions[r].height=62 if len(v)>150 else 34; r+=1
ws.column_dimensions["A"].width=22; ws.column_dimensions["B"].width=120

# ---------------- Raw ----------------
wr=wb.create_sheet("Raw_Data")
hdr(wr,1,1,list(raw.columns))
for row in raw.itertuples(index=False):
    wr.append([None if (isinstance(x,float) and np.isnan(x)) else x for x in row])
for i in range(1,raw.shape[1]+1): wr.column_dimensions[L(i)].width=16
wr.freeze_panes="A2"
for row in wr.iter_rows(min_row=2):
    for c in row: c.font=bfont

# ---------------- Cleaning log ----------------
wl=wb.create_sheet("Cleaning_Log"); wl.sheet_view.showGridLines=False
title(wl,1,"Data cleaning log","Each step, the Excel feature used, and the number of rows affected")
tool={"Rows in raw file":"-","Text columns trimmed & converted to Proper Case":"TRIM() + PROPER()  (or Power Query > Transform > Trim/Capitalize Each Word)",
"Discount values rescaled (20 -> 0.20)":"=IF(Q2>1,Q2/100,Q2)","Missing City filled from customer's other orders":"INDEX/MATCH on Customer ID",
"Missing City still unresolved -> 'Unknown'":"Find & Replace blanks","Rows removed: Ship Date before Order Date":"Filter on =C2<B2 then delete rows",
"Rows removed: Quantity <= 0":"Filter Quantity <= 0 then delete rows","Duplicate order lines removed":"Data > Remove Duplicates on Order ID + Product ID","Rows in clean file":"-"}
hdr(wl,4,1,["Step","Rows affected","Excel technique"])
for i,(s,n) in enumerate(log.itertuples(index=False),5):
    wl.cell(i,1,s).font=bfont; c=wl.cell(i,2,int(n)); c.font=Font(name=F,size=10,color="0000FF"); c.number_format="#,##0"; wl.cell(i,3,tool.get(s,"")).font=bfont
wl.cell(15,1,"Text dates in the raw file (dd-mm-yyyy) were converted to true Excel dates with DATEVALUE / Text-to-Columns.").font=Font(name=F,size=9,italic=True)
wl.column_dimensions["A"].width=55; wl.column_dimensions["B"].width=16; wl.column_dimensions["C"].width=70

# ---------------- Clean data ----------------
wc=wb.create_sheet("Clean_Data")
cols=list(clean.columns)+["Order Year","Order Month","Delivery Days","Profit Margin %","Discount Band","Loss Flag","New Order Flag"]
hdr(wc,1,1,cols)
for i,row in enumerate(clean.itertuples(index=False),2):
    vals=list(row); vals[1]=vals[1].to_pydatetime(); vals[2]=vals[2].to_pydatetime()
    for j,v in enumerate(vals,1): wc.cell(i,j,v).font=bfont
    f=[f"=YEAR(B{i})",f"=MONTH(B{i})",f"=C{i}-B{i}",f"=IF(O{i}=0,0,R{i}/O{i})",
       f'=IF(Q{i}=0,"0. No discount",IF(Q{i}<=0.1,"1. 1-10%",IF(Q{i}<=0.2,"2. 11-20%",IF(Q{i}<=0.3,"3. 21-30%","4. Above 30%"))))',
       f"=IF(R{i}<0,1,0)",f'=IF(A{i}<>A{i-1},1,0)' if i>2 else "=1"]
    for j,v in enumerate(f,19): wc.cell(i,j,v).font=bfont
    wc.cell(i,2).number_format="dd-mmm-yyyy"; wc.cell(i,3).number_format="dd-mmm-yyyy"
    wc.cell(i,15).number_format="#,##0.00"; wc.cell(i,18).number_format="#,##0.00"; wc.cell(i,17).number_format="0%"; wc.cell(i,22).number_format="0.0%"
for j in range(1,len(cols)+1): wc.column_dimensions[L(j)].width=15
wc.column_dimensions["N"].width=24; wc.column_dimensions["F"].width=20
wc.freeze_panes="C2"
t=Table(displayName="tblSales",ref=f"A1:{L(len(cols))}{last}"); t.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True); wc.add_table(t)
names={"rngSales":"O","rngQty":"P","rngDisc":"Q","rngProfit":"R","rngYear":"S","rngMonth":"T","rngDeliv":"U","rngBand":"W","rngLoss":"X","rngNewOrd":"Y",
 "rngCat":"L","rngSub":"M","rngRegion":"H","rngSegment":"G","rngShip":"D","rngCust":"E"}
for n,c in names.items(): wb.defined_names[n]=DefinedName(n,attr_text=f"Clean_Data!${c}$2:${c}${last}")

# ---------------- Customers ----------------
wu=wb.create_sheet("Customers")
cu=clean.groupby("Customer ID").agg(name=("Customer Name","first"),seg=("Segment","first"),reg=("Region","first")).reset_index()
hdr(wu,1,1,["Customer ID","Customer Name","Segment","Region","Sales","Profit","Orders","Profit Rank"])
nc=len(cu)
for i,r_ in enumerate(cu.itertuples(index=False),2):
    for j,v in enumerate(r_,1): wu.cell(i,j,v).font=bfont
    wu.cell(i,5,f"=SUMIFS(rngSales,rngCust,A{i})").number_format="#,##0"
    wu.cell(i,6,f"=SUMIFS(rngProfit,rngCust,A{i})").number_format="#,##0"
    wu.cell(i,7,f"=SUMIFS(rngNewOrd,rngCust,A{i})")
    wu.cell(i,8,f"=RANK(F{i},$F$2:$F${nc+1})+COUNTIF($F$2:F{i},F{i})-1")
    for j in (5,6,7,8): wu.cell(i,j).font=bfont
for j,w in enumerate([14,22,14,12,14,14,10,12],1): wu.column_dimensions[L(j)].width=w
wu.freeze_panes="A2"

# ---------------- Summary ----------------
S=wb.create_sheet("Summary"); S.sheet_view.showGridLines=False
S.column_dimensions["A"].width=26
for c in "BCDEFGH": S.column_dimensions[c].width=16
pos={}
def money(c): c.number_format='#,##0;[Red](#,##0);-'
def pct(c): c.number_format='0.0%;[Red]-0.0%;-'

title(S,1,"SUMMARY ANALYSIS - all figures are live formulas on Clean_Data","INR. Margin = Profit / Sales.")
# KPIs
r=4; hdr(S,r,1,["KPI","Value"])
kp=[("Total Sales","=SUM(rngSales)",money),("Total Profit","=SUM(rngProfit)",money),("Profit Margin","=B6/B5",pct),
    ("Orders","=SUM(rngNewOrd)",lambda c:setattr(c,"number_format","#,##0")),("Units Sold","=SUM(rngQty)",lambda c:setattr(c,"number_format","#,##0")),
    ("Avg Order Value","=B5/B8",money),("% Order Lines Loss-making","=SUM(rngLoss)/COUNT(rngSales)",pct),
    ("Avg Delivery Days (per order)","=AVERAGEIFS(rngDeliv,rngNewOrd,1)",lambda c:setattr(c,"number_format","0.0"))]
for i,(k,f,fm) in enumerate(kp,5):
    S.cell(i,1,k).font=bfont; c=S.cell(i,2,f); c.font=bold; fm(c)
pos["kpi"]=(5,12)

def block(r0,ttl,sub,heads,keys,formulas,fmts,keycol_fmt=None):
    title(S,r0,ttl,sub); hdr(S,r0+2,1,heads); r=r0+3
    for k in keys:
        c=S.cell(r,1,k); c.font=Font(name=F,size=10,color="0000FF") ; 
        for j,(f,fm) in enumerate(zip(formulas,fmts),2):
            x=S.cell(r,j,f.replace("{r}",str(r)).replace("{k}",f"$A{r}")); x.font=bfont; fm(x)
        r+=1
    return r0+3,r-1

# By year
yr0,yr1=block(15,"1. Sales & profit by year","YoY = change vs previous year",["Year","Sales","Profit","Margin","Sales YoY"],[2023,2024,2025],
  ["=SUMIFS(rngSales,rngYear,{k})","=SUMIFS(rngProfit,rngYear,{k})","=C{r}/B{r}","=IF(ROW()={first},\"-\",B{r}/B{rp}-1)"],[money,money,pct,pct])
# fix YoY formulas
for r in range(yr0,yr1+1):
    S.cell(r,5,"-" if r==yr0 else f"=B{r}/B{r-1}-1"); pct(S.cell(r,5)); S.cell(r,5).font=bfont
S.cell(yr0,1).number_format="0"
# by category
c0,c1=block(22,"2. Category performance","",["Category","Sales","Profit","Margin","Sales share"],["Furniture","Office Supplies","Technology"],
  ["=SUMIFS(rngSales,rngCat,{k})","=SUMIFS(rngProfit,rngCat,{k})","=C{r}/B{r}","=B{r}/SUM(rngSales)"],[money,money,pct,pct])
# by sub category ordered by profit
sub=clean.groupby("Sub-Category").Profit.sum().sort_values(ascending=False).index.tolist()
s0,s1=block(28,"3. Sub-category performance (sorted by profit at build time)","Red = loss-making",["Sub-Category","Sales","Profit","Margin","Avg Discount","Status"],sub,
  ["=SUMIFS(rngSales,rngSub,{k})","=SUMIFS(rngProfit,rngSub,{k})","=C{r}/B{r}","=AVERAGEIFS(rngDisc,rngSub,{k})","=IF(C{r}<0,\"LOSS-MAKING\",\"OK\")"],[money,money,pct,pct,lambda c:None])
S.conditional_formatting.add(f"C{s0}:C{s1}",CellIsRule(operator="lessThan",formula=["0"],fill=PatternFill("solid",bgColor="F8D7DA")))
S.conditional_formatting.add(f"F{s0}:F{s1}",CellIsRule(operator="equal",formula=['"LOSS-MAKING"'],fill=PatternFill("solid",bgColor="F8D7DA"),font=Font(color="9C0006",bold=True)))
# region
rg0,rg1=block(s1+3,"4. Region performance","",["Region","Sales","Profit","Margin","Avg Discount","% Lines Loss-making"],["South","West","North","East","Central"],
  ["=SUMIFS(rngSales,rngRegion,{k})","=SUMIFS(rngProfit,rngRegion,{k})","=C{r}/B{r}","=AVERAGEIFS(rngDisc,rngRegion,{k})","=SUMIFS(rngLoss,rngRegion,{k})/COUNTIFS(rngRegion,{k})"],[money,money,pct,pct,pct])
# segment
sg0,sg1=block(rg1+3,"5. Customer segment performance","",["Segment","Sales","Profit","Margin","Orders"],["Consumer","Corporate","Home Office"],
  ["=SUMIFS(rngSales,rngSegment,{k})","=SUMIFS(rngProfit,rngSegment,{k})","=C{r}/B{r}","=SUMIFS(rngNewOrd,rngSegment,{k})"],[money,money,pct,lambda c:setattr(c,"number_format","#,##0")])
# discount band
bands=["0. No discount","1. 1-10%","2. 11-20%","3. 21-30%","4. Above 30%"]
d0,d1=block(sg1+3,"6. Discount band analysis (key insight)","Margin turns negative above 20% discount",["Discount Band","Order Lines","Sales","Profit","Margin","% Lines Loss-making"],bands,
  ["=COUNTIFS(rngBand,{k})","=SUMIFS(rngSales,rngBand,{k})","=SUMIFS(rngProfit,rngBand,{k})","=D{r}/C{r}","=SUMIFS(rngLoss,rngBand,{k})/B{r}"],[lambda c:setattr(c,"number_format","#,##0"),money,money,pct,pct])
S.conditional_formatting.add(f"D{d0}:D{d1}",CellIsRule(operator="lessThan",formula=["0"],fill=PatternFill("solid",bgColor="F8D7DA")))
# ship mode
sm0,sm1=block(d1+3,"7. Shipping performance","",["Ship Mode","Orders","Avg Delivery Days"],["Same Day","First Class","Second Class","Standard Class"],
  ["=SUMIFS(rngNewOrd,rngShip,{k})","=AVERAGEIFS(rngDeliv,rngShip,{k},rngNewOrd,1)"],[lambda c:setattr(c,"number_format","#,##0"),lambda c:setattr(c,"number_format","0.00")])
# Standard class delivery by region
sr0,sr1=block(sm1+3,"8. Standard Class delivery days by region","Central lags other regions",["Region","Avg Delivery Days","% Orders > 7 days"],["Central","East","South","West","North"],
  ["=AVERAGEIFS(rngDeliv,rngRegion,{k},rngShip,\"Standard Class\",rngNewOrd,1)","=COUNTIFS(rngRegion,{k},rngShip,\"Standard Class\",rngNewOrd,1,rngDeliv,\">7\")/COUNTIFS(rngRegion,{k},rngShip,\"Standard Class\",rngNewOrd,1)"],[lambda c:setattr(c,"number_format","0.00"),pct])
# monthly matrix
title(S,sr1+3,"9. Monthly sales by year + seasonality index","Index 100 = average month")
mr=sr1+5; hdr(S,mr,1,["Month","2023","2024","2025","Total","Seasonality Index"])
mnames=["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
for m in range(1,13):
    r=mr+m; S.cell(r,1,mnames[m-1]).font=bfont
    for j,y in enumerate([2023,2024,2025],2):
        c=S.cell(r,j,f"=SUMIFS(rngSales,rngYear,{y},rngMonth,{m})"); money(c); c.font=bfont
    c=S.cell(r,5,f"=SUM(B{r}:D{r})"); money(c); c.font=bfont
    c=S.cell(r,6,f"=E{r}/AVERAGE($E${mr+1}:$E${mr+12})*100"); c.number_format="0"; c.font=bfont
mo0,mo1=mr+1,mr+12
S.conditional_formatting.add(f"F{mo0}:F{mo1}",DataBarRule(start_type="num",start_value=0,end_type="num",end_value=160,color=TEAL))
# top 10 customers
tr=mo1+3; title(S,tr,"10. Top 10 customers by profit","Dynamic: INDEX/MATCH on the rank column of the Customers sheet")
hdr(S,tr+2,1,["Rank","Customer ID","Customer Name","Profit","Cumulative % of total profit"])
for k in range(1,11):
    r=tr+2+k; S.cell(r,1,k).font=bfont
    S.cell(r,2,f"=INDEX(Customers!$A$2:$A${nc+1},MATCH(A{r},Customers!$H$2:$H${nc+1},0))")
    S.cell(r,3,f"=INDEX(Customers!$B$2:$B${nc+1},MATCH(A{r},Customers!$H$2:$H${nc+1},0))")
    c=S.cell(r,4,f"=INDEX(Customers!$F$2:$F${nc+1},MATCH(A{r},Customers!$H$2:$H${nc+1},0))"); money(c)
    c=S.cell(r,5,f"=SUM($D${tr+3}:D{r})/$B$6"); pct(c)
    for j in (2,3,4,5): S.cell(r,j).font=bfont
S.freeze_panes="A4"

# ---------------- Dashboard ----------------
D=wb.create_sheet("Dashboard",1); D.sheet_view.showGridLines=False
for i in range(1,25): D.column_dimensions[L(i)].width=11
D.merge_cells("A1:X2"); D["A1"]="RETAIL SALES PERFORMANCE DASHBOARD  |  Jan 2023 - Dec 2025  |  INR"
D["A1"].font=Font(name=F,size=16,bold=True,color="FFFFFF"); D["A1"].fill=hfill; D["A1"].alignment=Alignment(horizontal="left",vertical="center",indent=1)
cards=[("TOTAL SALES","=Summary!B5",'#,##0.0,,"M"'),("TOTAL PROFIT","=Summary!B6",'#,##0.0,,"M"'),("PROFIT MARGIN","=Summary!B7","0.0%"),
       ("ORDERS","=Summary!B8","#,##0"),("AVG ORDER VALUE","=Summary!B10","#,##0"),("% LOSS-MAKING LINES","=Summary!B11","0.0%")]
for i,(lab,f,fm) in enumerate(cards):
    c0=1+i*4
    D.merge_cells(start_row=4,start_column=c0,end_row=4,end_column=c0+3); D.merge_cells(start_row=5,start_column=c0,end_row=6,end_column=c0+3)
    a=D.cell(4,c0,lab); a.font=Font(name=F,size=9,bold=True,color="666666"); a.alignment=Alignment(horizontal="center"); 
    b=D.cell(5,c0,f); b.font=Font(name=F,size=20,bold=True,color=NAVY); b.number_format=fm; b.alignment=Alignment(horizontal="center",vertical="center")
    for rr in (4,5,6):
        for cc in range(c0,c0+4): D.cell(rr,cc).fill=PatternFill("solid",fgColor=LIGHT)

def style(ch,t,w=15.5,h=7.5):
    ch.title=t; ch.width=w; ch.height=h; ch.legend.position="b" if ch.legend else None
# chart 1 sales & profit by year
c1=BarChart(); c1.type="col"; c1.add_data(Reference(S,min_col=2,max_col=3,min_row=yr0-1,max_row=yr1),titles_from_data=True); c1.set_categories(Reference(S,min_col=1,min_row=yr0,max_row=yr1)); style(c1,"Sales & Profit by Year (INR)")
D.add_chart(c1,"A8")
# chart 2 monthly sales
c2=LineChart(); c2.add_data(Reference(S,min_col=2,max_col=4,min_row=mr,max_row=mr+12),titles_from_data=True); c2.set_categories(Reference(S,min_col=1,min_row=mo0,max_row=mo1)); style(c2,"Monthly Sales by Year - festive Q4 peak")
D.add_chart(c2,"M8")
# chart 3 subcategory profit
c3=BarChart(); c3.type="bar"; c3.add_data(Reference(S,min_col=3,min_row=s0-1,max_row=s1),titles_from_data=True); c3.set_categories(Reference(S,min_col=1,min_row=s0,max_row=s1)); style(c3,"Profit by Sub-Category (Tables & Machines lose money)",15.5,9); c3.legend=None
c3.y_axis.delete=False; c3.x_axis.delete=False; c3.x_axis.scaling.orientation="maxMin"; c3.x_axis.tickLblPos='low'
D.add_chart(c3,"A24")
# chart 4 discount band margin
c4=BarChart()
c4_tmp=None; c4.type="col"; c4.add_data(Reference(S,min_col=5,min_row=d0-1,max_row=d1),titles_from_data=True); c4.set_categories(Reference(S,min_col=1,min_row=d0,max_row=d1)); style(c4,"Profit Margin by Discount Band",15.5,9); c4.legend=None
c4.dataLabels=DataLabelList(); c4.dataLabels.showVal=True; c4.dataLabels.showSerName=False; c4.dataLabels.showCatName=False; c4.dataLabels.showLegendKey=False; c4.y_axis.number_format='0%'
c4.x_axis.tickLblPos='low'
D.add_chart(c4,"M24")
# chart 5 region margin
c5=BarChart(); c5.type="col"; c5.add_data(Reference(S,min_col=4,min_row=rg0-1,max_row=rg1),titles_from_data=True); c5.set_categories(Reference(S,min_col=1,min_row=rg0,max_row=rg1)); style(c5,"Profit Margin by Region",15.5,7.5); c5.legend=None
c5.dataLabels=DataLabelList(); c5.dataLabels.showVal=True; c5.dataLabels.showSerName=False; c5.dataLabels.showCatName=False; c5.dataLabels.showLegendKey=False; c5.y_axis.number_format='0%'
D.add_chart(c5,"A43")
# chart 6 seasonality
c6=BarChart(); c6.type="col"; c6.add_data(Reference(S,min_col=6,min_row=mr,max_row=mr+12),titles_from_data=True); c6.set_categories(Reference(S,min_col=1,min_row=mo0,max_row=mo1)); style(c6,"Seasonality Index (100 = average month)",15.5,7.5); c6.legend=None
D.add_chart(c6,"M43")
for sr in c2.series: sr.smooth=False
for ch in (c1,c2,c3,c4,c5,c6):
    ch.y_axis.delete=False; ch.x_axis.delete=False
D["A60"]="KEY INSIGHTS"; D["A60"].font=Font(name=F,size=11,bold=True,color=NAVY)
ins=["1) Margin collapses once discounts exceed 20%: lines above 30% discount lose money about 79% of the time.",
"2) Tables and Machines are loss-making overall - both are discounted heavily (avg ~21-22%).",
"3) Central has the lowest margin, the deepest discounts and the slowest Standard Class delivery.",
"4) Oct-Dec carries the festive-season peak (seasonality index 121-137) - plan stock and discount caps for it."]
for k,t in enumerate(ins): D.cell(61+k,1,t).font=bfont
D.page_setup.orientation="landscape"; D.page_setup.fitToWidth=1; D.page_setup.fitToHeight=0; D.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
wb.save("excel/Retail_Sales_Analysis.xlsx"); print("saved",S.max_row)
json.dump(dict(s0=s0,s1=s1),open("data/_layout.json","w"))
